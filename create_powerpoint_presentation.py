import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette: Clean White & Royal/Navy Blue
COLOR_BG = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(248, 250, 252) # Slate 50
COLOR_CARD_BORDER = RGBColor(203, 213, 225) # Slate 300
COLOR_PRIMARY_NAVY = RGBColor(15, 43, 92) # #0f2b5c
COLOR_ROYAL_BLUE = RGBColor(30, 64, 175) # #1e40af
COLOR_ACCENT_BLUE = RGBColor(2, 132, 199) # #0284c7
COLOR_TEXT_DARK = RGBColor(15, 23, 42) # Slate 900
COLOR_TEXT_MUTED = RGBColor(71, 85, 105) # Slate 600
COLOR_WHITE = RGBColor(255, 255, 255)

COLOR_GREEN_BG = RGBColor(236, 253, 245)
COLOR_GREEN_TEXT = RGBColor(5, 150, 105)
COLOR_AMBER_BG = RGBColor(255, 251, 235)
COLOR_AMBER_TEXT = RGBColor(217, 119, 6)
COLOR_RED_BG = RGBColor(254, 242, 242)
COLOR_RED_TEXT = RGBColor(220, 38, 38)

FONT_ARABIC = "Cairo"
FONT_MONO = "Consolas"

blank_slide_layout = prs.slide_layouts[6]

def add_header(slide, slide_num, category, title, subtitle, status_text=None, status_type="blue"):
    # Top accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ROYAL_BLUE
    bar.line.fill.background()

    # Category & Slide Counter
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(8.5), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"• {category} | مشروع LOCAL CONTACT P2P"
    p_cat.font.name = FONT_ARABIC
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_ACCENT_BLUE

    # Slide Number pill
    num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.5), Inches(0.25), Inches(1.0), Inches(0.35))
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
    num_box.line.fill.background()
    p_num = num_box.text_frame.paragraphs[0]
    p_num.text = f"{slide_num} / 12"
    p_num.alignment = PP_ALIGN.CENTER
    p_num.font.name = FONT_MONO
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_WHITE

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(9.5), Inches(0.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_ARABIC
    p_t.font.size = Pt(20)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_PRIMARY_NAVY

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.18), Inches(9.5), Inches(0.4))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle
    p_sub.font.name = FONT_ARABIC
    p_sub.font.size = Pt(11.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Status Pill if any
    if status_text:
        bg_col = COLOR_ROYAL_BLUE
        txt_col = COLOR_WHITE
        if status_type == "green":
            bg_col = COLOR_GREEN_TEXT
        elif status_type == "amber":
            bg_col = COLOR_AMBER_TEXT
        elif status_type == "red":
            bg_col = COLOR_RED_TEXT
        
        stat_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), Inches(0.75), Inches(2.5), Inches(0.38))
        stat_shape.fill.solid()
        stat_shape.fill.fore_color.rgb = bg_col
        stat_shape.line.fill.background()
        p_stat = stat_shape.text_frame.paragraphs[0]
        p_stat.text = status_text
        p_stat.alignment = PP_ALIGN.CENTER
        p_stat.font.name = FONT_MONO
        p_stat.font.size = Pt(10)
        p_stat.font.bold = True
        p_stat.font.color.rgb = txt_col

    # Bottom line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.73), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.fill.background()

def set_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "🎙️ ملاحظات المتحدث الشفهية للمناقشة:\n" + notes_text

# ==============================================================================
# SLIDE 1: Title & Cover
# ==============================================================================
s1 = prs.slides.add_slide(blank_slide_layout)

# Decorative top & bottom blue bands
top_band = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
top_band.fill.solid()
top_band.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
top_band.line.fill.background()

bottom_band = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.3), Inches(13.333), Inches(0.2))
bottom_band.fill.solid()
bottom_band.fill.fore_color.rgb = COLOR_ROYAL_BLUE
bottom_band.line.fill.background()

# Title badge
badge_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.8), Inches(0.5), Inches(5.7), Inches(0.45))
badge_box.fill.solid()
badge_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
badge_box.line.color.rgb = COLOR_ROYAL_BLUE
badge_p = badge_box.text_frame.paragraphs[0]
badge_p.text = "🛡️ المناقشة الأكاديمية — هندسة البرمجيات والتشفير المتقدم"
badge_p.alignment = PP_ALIGN.CENTER
badge_p.font.name = FONT_ARABIC
badge_p.font.size = Pt(12)
badge_p.font.bold = True
badge_p.font.color.rgb = COLOR_ROYAL_BLUE

# Main Title
title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.7), Inches(0.8))
tf_main = title_box.text_frame
tf_main.word_wrap = True
p_main = tf_main.paragraphs[0]
p_main.text = "مناقشة خوارزميات التشفير في مشروع LOCAL CONTACT"
p_main.alignment = PP_ALIGN.CENTER
p_main.font.name = FONT_ARABIC
p_main.font.size = Pt(28)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_PRIMARY_NAVY

# Subtitle
sub_box = s1.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(10.3), Inches(0.5))
tf_s = sub_box.text_frame
tf_s.word_wrap = True
p_s = tf_s.paragraphs[0]
p_s.text = "التصميم الأمني وهندسة خوارزميات التشفير في نظام اتصال P2P محلي مستقل يعمل دون خادم مركزي"
p_s.alignment = PP_ALIGN.CENTER
p_s.font.name = FONT_ARABIC
p_s.font.size = Pt(13)
p_s.font.bold = True
p_s.font.color.rgb = COLOR_TEXT_MUTED

# 6 Cryptographic Algorithms Grid (White cards with blue borders)
algos = [
    ("AES-256-GCM", "التشفير المتماثل ومصادقة البيانات (AEAD)"),
    ("ECDH (NIST P-256)", "تبادل المفاتيح وتأمين Forward Secrecy"),
    ("HKDF-SHA256", "دالة اشتقاق مفاتيح الجلسات (RFC 5869)"),
    ("ECDSA + Keystore", "التوقيع الرقمي وهوية الجهاز بالمعالج الآمن"),
    ("SHA-256 Fingerprint", "بصمات الأجهزة ونموذج الثقة (TOFU)"),
    ("PBKDF2-HMAC-SHA256", "المفتاح الشبكي الموحد للتوافق العكسي")
]

for idx, (aname, arole) in enumerate(algos):
    col = idx % 3
    row = idx // 3
    x = Inches(0.8 + col * 3.95)
    y = Inches(2.45 + row * 1.1)
    
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(0.95))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_ROYAL_BLUE
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = aname
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = FONT_MONO
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_ROYAL_BLUE
    
    p2 = tf.add_paragraph()
    p2.text = arole
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_ARABIC
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_TEXT_DARK

# Students Box (Names in Blue on White)
team_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.75), Inches(11.7), Inches(1.3))
team_box.fill.solid()
team_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
team_box.line.color.rgb = COLOR_ROYAL_BLUE
team_box.line.width = Pt(1.5)

tf_team = team_box.text_frame
p_tt = tf_team.paragraphs[0]
p_tt.text = "👥 الطلاب المنفذون للمشروع:"
p_tt.alignment = PP_ALIGN.CENTER
p_tt.font.name = FONT_ARABIC
p_tt.font.size = Pt(12)
p_tt.font.bold = True
p_tt.font.color.rgb = COLOR_PRIMARY_NAVY

students = ["مشعل حاجب", "قحطان الشاجع", "أواب النزيلي", "محمد العيدروس", "محمد العواضي"]
for sidx, sname in enumerate(students):
    sx = Inches(1.1 + sidx * 2.25)
    sy = Inches(5.35)
    scard = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, Inches(2.1), Inches(0.5))
    scard.fill.solid()
    scard.fill.fore_color.rgb = COLOR_WHITE
    scard.line.color.rgb = COLOR_ACCENT_BLUE
    scard.line.width = Pt(1.2)
    sp = scard.text_frame.paragraphs[0]
    sp.text = f"🎓 {sname}"
    sp.alignment = PP_ALIGN.CENTER
    sp.font.name = FONT_ARABIC
    sp.font.size = Pt(11)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_ROYAL_BLUE

# Footer Info
foot_box = s1.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
tf_f = foot_box.text_frame
p_f = tf_f.paragraphs[0]
p_f.text = "لجنة المناقشة: أستاذ التشفير & أستاذ هندسة البرمجيات  |  الحالة: BUILD SUCCESSFUL  |  الاختبارات: 33 / 33 ناجح على JVM"
p_f.alignment = PP_ALIGN.CENTER
p_f.font.name = FONT_ARABIC
p_f.font.size = Pt(11)
p_f.font.bold = True
p_f.font.color.rgb = COLOR_TEXT_MUTED

set_speaker_notes(s1, "اشرح للأساتذة في الافتتاحية أن هذا العرض يركز على مناقشة خوارزميات التشفير في مشروع LOCAL CONTACT، وهو نظام اتصال محلي P2P يعمل دون خادم مركزي. وقدم زملاءك في فريق العمل: مشعل حاجب، قحطان الشاجع، أواب النزيلي، محمد العيدروس، ومحمد العواضي، مع استعراض الخوارزميات الأساسية الست المستخدمة في النظام.")

# ==============================================================================
# Helper to create content cards
# ==============================================================================
def add_card(slide, x, y, w, h, title, bullets, file_info=None, border_color=COLOR_ROYAL_BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_ARABIC
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_NAVY
    
    for b in bullets:
        pb = tf.add_paragraph()
        pb.text = f"• {b}"
        pb.font.name = FONT_ARABIC
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = COLOR_TEXT_DARK
        pb.space_before = Pt(4)
        
    if file_info:
        pf_title = tf.add_paragraph()
        pf_title.text = "📂 المكون البرمجي المرتبط:"
        pf_title.font.name = FONT_ARABIC
        pf_title.font.size = Pt(10.5)
        pf_title.font.bold = True
        pf_title.font.color.rgb = COLOR_ROYAL_BLUE
        pf_title.space_before = Pt(8)
        
        pf_path = tf.add_paragraph()
        pf_path.text = file_info.get("path", "")
        pf_path.font.name = FONT_MONO
        pf_path.font.size = Pt(9.5)
        pf_path.font.bold = True
        pf_path.font.color.rgb = COLOR_ACCENT_BLUE
        
        pf_desc = tf.add_paragraph()
        pf_desc.text = file_info.get("desc", "")
        pf_desc.font.name = FONT_ARABIC
        pf_desc.font.size = Pt(9.5)
        pf_desc.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# SLIDE 2: What was the problem?
# ==============================================================================
s2 = prs.slides.add_slide(blank_slide_layout)
add_header(s2, 2, "نقطة الانطلاق والتحليل المعماري", "ما المشكلة الأمنية في التصميم السابق؟", "تشخيص مكامن الضعف في إدارة المفاتيح وحماية الجلسات قبل التحسين", "CRITICAL GAP", "red")

add_card(s2, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "⚠️ المخاطر ونقاط الضعف المعمارية السابقة:",
    [
        "مفتاح شبكي موحد ثابت: اشتقاق مفتاح AES-256 عبر PBKDF2 من عبارة مرور ثابتة داخل الكود.",
        "مفتاح مشترك بين الأجهزة والجلسات: فك تشفير حزمة واحدة يكشف بيانات باقي الأجهزة.",
        "غياب تبادل المفاتيح الديناميكي: عدم وجود أي بروتوكول لتبادل المفاتيح بين الأجهزة.",
        "انعدام سرية التوجيه المستقبلي: كشف عبارة المرور يفك تشفير كل الاتصالات السابقة المسجلة.",
        "غياب الهوية المشفرة: الاعتماد على معرفات عشوائية دون مفاتيح غير متماثلة.",
        "الخلاصة: خوارزمية AES-GCM قوية، لكن طريقة إدارة وتوزيع المفتاح كانت نقطة الضعف."
    ],
    border_color=COLOR_RED_TEXT
)

add_card(s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 الملف والمسؤولية البرمجية المعنية:",
    [
        "اقتران التشفير مع المفتاح: كان محرك التشفير يقوم بتشفير الرسائل واشتقاق المفتاح الثابت داخلياً بنفسه.",
        "صعوبة التبديل والتطوير: أي رغبة في استخدام مفاتيح مختلفة كانت تتطلب إعادة كتابة دوال التشفير بالكامل.",
        "غياب التوثيق والتتبع: عدم وجود مواصفات متطلبات تربط الكود بمرجعية أمنية معتمدة."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/LocalCryptoEngine.kt",
        "desc": "المسؤولية السابقة: محرك التشفير المتماثل واشتقاق المفتاح الثابت من عبارة المرور الموحدة."
    }
)
set_speaker_notes(s2, "وضّح للدكاترة الفجوة الأمنية السابقة: التطبيق كان يحتوي على وظائف اتصال ممتازة ولكن أمنياً كان يعتمد على عبارة مرور ثابتة مشتقة عبر PBKDF2، مما جعل مفتاح التشفير ثابتاً ومشتركاً بين جميع الأجهزة وكافة الجلسات. كشف المفتاح من الـ APK كان يهدد بفك تشفير كل الحزم المسجلة.")

# ==============================================================================
# SLIDE 3: KeyProvider
# ==============================================================================
s3 = prs.slides.add_slide(blank_slide_layout)
add_header(s3, 3, "هندسة البرمجيات والتصميم المعماري", "فصل إدارة المفاتيح عبر تجريد KeyProvider", "تطبيق مبادئ SOLID ونمط Provider لعزل مسؤولية اشتقاق المفاتيح", "IMPLEMENTED", "green")

add_card(s3, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "🧩 المبادئ الهندسية ونمط التصميم المطبق:",
    [
        "تطبيق مبدأ عكس التبعيات (DIP): محرك التشفير لم يعد يشتق المفتاح بنفسه، بل يعتمد على واجهة مجردة.",
        "مبدأ المسؤولية الفردية (SRP): عزل مسؤولية توليد المفاتيح عن عمليات التشفير والـ Cipher.",
        "نمط المزود (Provider / Strategy Pattern): إمكانية إضافة استراتيجيات مفاتيح جديدة مستقبلاً دون المساس بمحرك التشفير.",
        "التوافق العكسي (Backward Compatibility): الحفاظ على سلوك المفتاح القديم مؤقتاً عبر LegacyStaticKeyProvider."
    ]
)

add_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 الملفات المضافة والمعدلة:",
    [
        "تجريد المفاتيح: إنشاء واجهة KeyProvider البرمجية.",
        "التوافق مع القديم: تنفيذ كلاس LegacyStaticKeyProvider للحفاظ على سلوك PBKDF2 القديم مؤقتاً.",
        "تعديل المحرك: تحديث LocalCryptoEngine ليعتمد على KeyProvider لتمرير المفتاح النشط."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/KeyProvider.kt",
        "desc": "ملف جديد يحتوي واجهة KeyProvider وتطبيق LegacyStaticKeyProvider لعزل التبعية."
    }
)
set_speaker_notes(s3, "هنا يظهر دور هندسة البرمجيات: طبقنا مبدأ Single Responsibility ومبدأ Dependency Inversion عبر إنشاء واجهة KeyProvider. عزلنا مسؤولية جلب المفتاح عن محرك التشفير. أنشأنا LegacyStaticKeyProvider للحفاظ على التوافق، ومهدنا الطريق لمفاتيح الجلسات الديناميكية دون كسر الكود القديم.")

# ==============================================================================
# SLIDE 4: AES-256-GCM
# ==============================================================================
s4 = prs.slides.add_slide(blank_slide_layout)
add_header(s4, 4, "التشفير المتماثل المعتمد", "التشفير المتماثل الأساسي: AES-256-GCM", "الحفاظ على التشفير المصادق المعتمد لضمان السرية وسلامة البيانات", "JVM TESTED (10/10)", "green")

add_card(s4, Inches(0.8), Inches(1.8), Inches(3.7), Inches(5.1),
    "🔒 السرية التامة:",
    [
        "مفتاح AES بطول 256 بت.",
        "خوارزمية قياسية معتمدة من NIST (FIPS 197).",
        "أداء سريع ومعالجة متوازية على الأجهزة المحمولة."
    ]
)

add_card(s4, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.1),
    "🛡️ سلامة البيانات (AEAD):",
    [
        "وسم مصادقة بطول 128 بت (Tag).",
        "رفض فك التشفير تلقائياً عند أي تعديل في النص المشفر.",
        "حماية كاملة من هجمات التلاعب بالبتات (Bit-Flipping)."
    ]
)

add_card(s4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.1),
    "🎲 العشوائية والـ Nonce:",
    [
        "رقم عشوائي بطول 12 بايت لكل حزمة.",
        "اشتقاق Nonce فريد لكل كتلة ملف (64KB).",
        "منع تكرار الـ Nonce تحت نفس المفتاح."
    ],
    file_info={
        "path": "app/src/test/.../LocalCryptoEngineTest.kt",
        "desc": "10 اختبارات ناجحة أثبتت كشف التلاعب ورفض المفتاح الخاطئ."
    }
)
set_speaker_notes(s4, "أكد لأستاذ التشفير أننا حافظنا على AES-256-GCM كخيار قياسي ممتاز لأنه نمط AEAD (يوفر السرية وسلامة البيانات معاً). وسيلة المصادقة Tag بطول 128 بت تكشف أي تلاعب بالبتات فوراً. اختبارات LocalCryptoEngineTest فحصت النصوص، كتل الملفات، والتلاعب وحققت 10/10 على JVM.")

# ==============================================================================
# SLIDE 5: ECDH P-256
# ==============================================================================
s5 = prs.slides.add_slide(blank_slide_layout)
add_header(s5, 5, "تبادل المفاتيح الديناميكي", "بروتوكول تبادل المفاتيح المنحني: ECDH (NIST P-256)", "توليد أسرار مشتركة مؤقتة لتحقيق سرية التوجيه المستقبلي (Forward Secrecy)", "JVM TESTED (8/8)", "green")

add_card(s5, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "⚡ الخصائص الأمنية لمنحنى NIST P-256:",
    [
        "المنحنى المعتمد: secp256r1 (prime256v1 / NIST P-256).",
        "مفاتيح مؤقتة (Ephemeral Keys): توليد زوج مفاتيح جديد لكل جلسة اتصال.",
        "إتلاف المفتاح الخاص: التخلص من المفتاح الخاص المؤقت فور حساب السر المشترك.",
        "سرية التوجيه المستقبلي (Forward Secrecy): اختراق الجهاز مستقبلاً لا يكشف الاتصالات القديمة المسجلة.",
        "توافق البيئة: مدعوم أصلياً في JCA داخل Android و JVM دون مكتبات خارجية ثقيلة."
    ]
)

add_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 المكون المنفذ والاختبارات:",
    [
        "توليد المفاتيح: توليد مفاتيح EC مؤقتة 256 بت.",
        "حساب السر المشترك: تطبيق KeyAgreement عبر ECDH.",
        "ترميز قياسي: تشفير المفاتيح العامة بصيغة X.509 DER القياسية.",
        "اختبارات الاتفاق: 8 اختبارات آلية أثبتت اتفاق طرفين مستقلين على نفس السر المشترك."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/EcdhEngine.kt",
        "desc": "محرك تبادل المفاتيح ECDH المنفذ والمختبر في EcdhEngineTest (8/8 ناجح)."
    }
)
set_speaker_notes(s5, "شرح بروتوكول تبادل المفاتيح: أنشأنا محرك ECDH باستخدام المنحنى الإهليلجي القياسي NIST P-256 (secp256r1). الميزة الأهم هي Forward Secrecy: الطرفان يولدّان مفاتيح مؤقتة Ephemeral لكل جلسة، يحسبان السر المشترك، ثم يُتلف المفتاح الخاص فوراً. لماذا P-256؟ لأنه مدعوم أصلياً في مكتبات التشفير القياسية داخل Android و JVM دون الحاجة لمكتبات خارجية ثقيلة.")

# ==============================================================================
# SLIDE 6: HKDF-SHA256
# ==============================================================================
s6 = prs.slides.add_slide(blank_slide_layout)
add_header(s6, 6, "دوال اشتقاق المفاتيح القياسية", "اشتقاق مفاتيح الجلسات عبر HKDF-SHA256", "تطبيق معيار RFC 5869 لتحويل السر المشترك إلى مفتاح متماثل نقي عشوائياً", "RFC 5869 COMPLIANT", "green")

add_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "⚙️ آليات دالة الاشتقاق (Extract-and-Expand):",
    [
        "مرحلة الاستخلاص (Extract): استخلاص العشوائية من السر المشترك عبر HMAC-SHA256 لإنتاج Pseudorandom Key (PRK).",
        "مرحلة التوسيع (Expand): توليد مفتاح تشفير 256 بت نقي باستخدام مصفوفة معلومات سياق الجلسة (Session Context Info).",
        "إزالة الانحياز الإحصائي: السر الناتج عن ECDH يمثل نقطة على المنحنى ولا يمتلك توزيعاً عشوائياً منتظماً يناسب AES مباشرة؛ HKDF يعالج ذلك.",
        "عزل الجلسات: ربط الاشتقاق بمعرفات الجلسة يمنع استخدام نفس المفتاح في سياق آخر."
    ]
)

add_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 الموضع البرمجي وأهميته الأكاديمية:",
    [
        "المعيار الذهبي: HKDF هو المعيار المعتمد في بروتوكولات TLS 1.3 و Signal و IPsec.",
        "التكامل في الكود: مدمج مباشرة في كلاس EcdhEngine لإنتاج كائن SecretKey.",
        "الاختبارات: تم التحقق من سلامة الاشتقاق وثبات المفاتيح بين طرفين مستقلين."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/EcdhEngine.kt",
        "desc": "دالتا deriveSharedSecret و deriveSessionKey لتطبيق مراحل HKDF."
    }
)
set_speaker_notes(s6, "وضّح لأستاذ التشفير لماذا نحتاج HKDF: السر المشترك الناتج عن ECDH يمثل نقطة على المنحنى، ولا يمتلك توزيعاً عشوائياً منتظماً (Uniform Distribution) يناسب مفاتيح AES. معيار RFC 5869 يطبق مرحلتين: Extract لاستخلاص العشوائية، و Expand لتوليد مفتاح AES-256 نقي مع حقن سياق الجلسة كـ Context لمنع الخلط بين الجلسات.")

# ==============================================================================
# SLIDE 7: Android Keystore Identity
# ==============================================================================
s7 = prs.slides.add_slide(blank_slide_layout)
add_header(s7, 7, "هوية الجهاز التشفيرية", "حماية هوية الجهاز عبر Android Keystore", "إدارة مفاتيح غير قابلة للتصدير مدعومة بالمعالج الآمن والتوقيع الرقمي ECDSA", "INTEGRATED IN CODE", "amber")

add_card(s7, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "🛡️ الخصائص الأمنية لمخزن المفاتيح:",
    [
        "مفتاح الهوية الدائم: توليد زوج مفاتيح EC secp256r1 فريد للجهاز داخل AndroidKeyStore.",
        "حماية العتاد الآمن (TEE / StrongBox): المفتاح الخاص لا يغادر المعالج المعزول أبداً ولا يمكن تصديره حتى بصلاحيات Root.",
        "التوقيع الرقمي: استخدام SHA256withECDSA لتوقيع حزم المصافحة وإثبات الهوية.",
        "عزل النطاق: مفتاح الهوية مخصص للمصادقة فقط ولا يشفر بيانات التطبيق مباشرة."
    ]
)

add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 المكون المسؤول وحالة التحقق الأكاديمية:",
    [
        "المسؤولية: إدارة دورة حياة المفتاح وتوقيع الحزم والتحقق منها.",
        "الفائدة: منح الجهاز هوية تشفيرية ثابتة بدل الاعتماد على معرف UUID عشوائي.",
        "⚠️ حالة التحقق الصارمة: مدمج في الكود ولكن يحتاج فحصاً على جهاز فيزيائي متاح يعمل بـ TEE حقيقي (Device Verification Pending)."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/KeystoreIdentityManager.kt",
        "desc": "كلاس إدارة هوية الجهاز والتوقيع الرقمي عبر Android Keystore."
    },
    border_color=COLOR_AMBER_TEXT
)
set_speaker_notes(s7, "اشرح كيف حلت هوية الجهاز مشكلة عدم وجود خادم: أنشأنا KeystoreIdentityManager لتوليد زوج مفاتيح دائم في Android Keystore المحمي بالعتاد (TEE أو StrongBox). المفتاح الخاص لا يغادر العتاد أبداً ولا يمكن تصديره أو سرقته حتى لو تم عمل روت للجهاز. يستخدم المفتاح حصراً لتوقيع حزم المصافحة بـ ECDSA. بيّن بأمانة أن هذا المكون مدمج بالكود ومصمم للأجهزة، ولم يخضع لاختبار Android Instrumented على محاكي.")

# ==============================================================================
# SLIDE 8: Context-Bound Signature
# ==============================================================================
s8 = prs.slides.add_slide(blank_slide_layout)
add_header(s8, 8, "تحصين بروتوكول المصافحة", "التوقيع الرقمي المرتبط بالسياق (Context-Bound)", "ربط حزم المصافحة بسياق الجلسة عبر هيكل باينري حتمي لمنع هجمات الاستبدال", "JVM TESTED (8/8)", "green")

add_card(s8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📦 حقول التوقيع الحتمي (Canonical Binary Layout):",
    [
        "Domain Tag: تمييز المجال بسلسلة LocalConnect::Handshake::v1.",
        "Protocol Version: رقم إصدار البروتوكول (الإصدار 2).",
        "Sender & Receiver IDs: معرفات أطراف الاتصال المحددة.",
        "Session ID: معرف الجلسة الفريد لمنع الخلط بين الجلسات.",
        "Timestamp: الطابع الزمني بالملي ثانية.",
        "Ephemeral Public Key: بايتات المفتاح العام المؤقت للقرين."
    ]
)

add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 المكونات المسؤولة واختبارات الأمان:",
    [
        "منع الاستبدال: تغيير أي بايت في سياق الجلسة يبطل التوقيع فوراً.",
        "منع إعادة التوجيه: لا يمكن استخدام نفس التوقيع ضد طرف آخر.",
        "الاختبارات الآلية: 8 اختبارات أثبتت كشف التلاعب بالطابع الزمني أو الجلسة.",
        "تنبيه هندسي: هذا يحصن المصافحة لكنه لا يمثل Replay Protection كاملاً لكافة الحزم الشبكية."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/HandshakeCryptoUtils.kt",
        "desc": "أدوات بناء مصفوفة التوقيع الحتمية وحساب بصمات SHA-256."
    }
)
set_speaker_notes(s8, "شرح تحصين المصافحة: في التصميم البدائي كان التوقيع يغطي المفتاح المؤقت فقط، مما يسمح بهجمات إعادة الاستخدام في جلسة أخرى. قمنا بإنشاء Canonical Binary Format يحزم: إصدار البروتوكول، معرف المرسل، معرف المستقبل، معرف الجلسة، الطابع الزمني، والمفتاح المؤقت. أي تعديل في سياق الجلسة يكسر التوقيع الرقمي فوراً. نبه اللجنة بأن هذا يحمي من التلاعب بالسياق، لكنه ليس Replay Protection كاملاً لبيانات الشبكة.")

# ==============================================================================
# SLIDE 9: TOFU Trust Store
# ==============================================================================
s9 = prs.slides.add_slide(blank_slide_layout)
add_header(s9, 9, "إدارة الثقة اللامركزية", "تثبيت هوية الأقران بنموذج TOFU", "سياسة الثقة عند أول اتصال (Trust On First Use) ومصادقة البصمات الرقمية", "PARTIAL MITM", "amber")

add_card(s9, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "🤝 آلية عمل نموذج TOFU:",
    [
        "في أول اتصال: تسجيل وحفظ بصمة المفتاح العام (SHA-256 Fingerprint) بمخزن الثقة المحلي.",
        "في الاتصالات اللاحقة: مطابقة البصمة القادمة مع البصمة المحفوظة؛ قبول الجلسة فوراً عند التطابق.",
        "كشف التغير: في حال تغير مفتاح القرين، يُرفض الاتصال وتنبيه المستخدم لاحتمال وجود هجوم.",
        "إدارة الحالات: تتبع حالة الأمان (LEGACY, IN_PROGRESS, SECURE) لمنع هجمات التراجع (Downgrade)."
    ]
)

add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "📂 المكونات البرمجية والحدود الأمنية:",
    [
        "مخزن الثقة: تخزين بصمات الأقران ومقارنتها بدقة.",
        "إدارة الجلسات: ضبط مهلة المصافحة (30 ثانية) وإلغاء الجلسات المنتهية.",
        "⚠️ الاعتراف الأكاديمي الصارم: نموذج TOFU يحمي بامتياز بدءاً من الاتصال الثاني فصاعداً. حماية أول اتصال تظل جزئية لغياب التحقق اليدوي (SAS)."
    ],
    file_info={
        "path": "app/src/main/java/com/example/network/PeerTrustStore.kt",
        "desc": "كلاس مخزن الثقة المحلي وتنسيق الحالات في PeerSessionKeyManager.kt."
    },
    border_color=COLOR_AMBER_TEXT
)
set_speaker_notes(s9, "شرح نموذج TOFU: لعدم وجود خادم شهادات مركزي CA في الشبكة المحلية، نعتمد Trust On First Use. في أول اتصال، يُحفظ المفتاح العام وبصمته. في الاتصالات اللاحقة يُرفض أي تغيير. اعترف بشفافية أكاديمية أمام أستاذ التشفير: حماية MITM في الاتصال الأول تظل جزئية وليست كاملة، لأن المهاجم النشط في اللحظة الأولى يمكنه انتحال الهوية إذا لم يتأكد الطرفان بقناة خارجية (SAS أو QR Code).")

# ==============================================================================
# SLIDE 10: Key Scope Reality
# ==============================================================================
s10 = prs.slides.add_slide(blank_slide_layout)
add_header(s10, 10, "الشفافية الهندسية والواقع العملي", "نطاق المفاتيح الفعلي في التطبيق (Key Scope)", "توضيح دقيق للمسارات التي تم ترقيتها لمفاتيح الجلسات والمسارات المحتفظة بالمفتاح الشبكي", "PARTIAL SCOPE", "amber")

# Table in slide 10
rows = 6
cols = 4
top = Inches(1.8)
left = Inches(0.8)
width = Inches(11.73)
height = Inches(4.2)

tbl_shape = s10.shapes.add_table(rows, cols, left, top, width, height)
tbl = tbl_shape.table

# Set Column Widths
tbl.columns[0].width = Inches(3.0)
tbl.columns[1].width = Inches(3.2)
tbl.columns[2].width = Inches(2.3)
tbl.columns[3].width = Inches(3.23)

headers = ["نوع البيانات / المسار", "المفتاح المستخدم حالياً", "المنفذ والبروتوكول", "الحالة المعمارية والأمنية"]
for cidx, h in enumerate(headers):
    cell = tbl.cell(0, cidx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_ARABIC
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

table_data = [
    ("الدردشة المباشرة (Direct Chat)", "sessionKey (ديناميكي)", "8888 (UDP Unicast)", "محمي بجلسة ECDH ديناميكية لكل قرين"),
    ("استكشاف الأجهزة (Discovery)", "networkKey (مشترك)", "8888 (UDP Broadcast)", "طبيعي (مطلوب عاماً للاكتشاف الأولي)"),
    ("غرف الدردشة الجماعية (Rooms)", "networkKey (قديم)", "8888 (UDP Multicast)", "مسار قديم (يحتاج Group Key مستقبلاً)"),
    ("نقل وتبادل الملفات (File Transfer)", "networkKey (قديم)", "8891 (TCP Stream)", "مشفر بكتل 64KB بمفتاح الشبكة القديم"),
    ("المكالمات الصوتية والمرئية", "networkKey (قديم)", "8889 / 8890 (UDP)", "مشفر بحزم PCM/JPEG بمفتاح الشبكة")
]

for ridx, row in enumerate(table_data):
    for cidx, val in enumerate(row):
        cell = tbl.cell(ridx + 1, cidx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if ridx % 2 == 0 else COLOR_WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_ARABIC if cidx != 1 else FONT_MONO
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.alignment = PP_ALIGN.CENTER

# Summary note
s10_note = s10.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.73), Inches(0.8))
tf_s10 = s10_note.text_frame
p_s10 = tf_s10.paragraphs[0]
p_s10.text = "📂 الملف المسؤول عن الدمج: LocalP2PEngine.kt | الأمانة الأكاديمية تقتضي عدم الادعاء بأن جميع أنواع البيانات تستخدم Session Keys."
p_s10.font.name = FONT_ARABIC
p_s10.font.size = Pt(11)
p_s10.font.bold = True
p_s10.font.color.rgb = COLOR_ROYAL_BLUE

set_speaker_notes(s10, "هذه الشريحة تبرز النضج والأمانة الهندسية: لا ندعي أن كل شيء يستخدم Session Keys. الدردشة الثنائية المباشرة Direct Chat هي التي تستخدم مفتاح الجلسة ECDH. استكشاف الأجهزة Discovery يستخدم المفتاح الشبكي وهو أمر طبيعي ومطلوب للاكتشاف العام. أما الغرف ونقل الملفات والصوت والفيديو فما زالت تستخدم المفتاح الشبكي القديم المشترك وتحتاج مفاتيح مجموعات مستقبلاً.")

# ==============================================================================
# SLIDE 11: Test Evidence & Limits
# ==============================================================================
s11 = prs.slides.add_slide(blank_slide_layout)
add_header(s11, 11, "إثباتات الجودة ومصفوفة التتبع", "ما الذي تثبته الاختبارات وما الذي لا تثبته؟", "تحليل نتيجة 33/33 اختباراً فحصاً برمجياً على JVM والحدود الأكاديمية الصارمة", "33/33 PASSED", "green")

add_card(s11, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "✅ ما تم إثباته برمجياً (33 اختباراً ناجحاً على JVM):",
    [
        "LocalCryptoEngineTest (10/10): صحة AES-GCM، كشف التلاعب بالنص المشفر والـ Tag، رفض المفتاح الخاطئ، ودعم JSON و UTF-8.",
        "EcdhEngineTest (8/8): اتفاق طرفين مستقلين على نفس السر المشترك وتوليد مفاتيح متطابقة عبر HKDF.",
        "PeerSessionKeyManagerTest (4/4): صحة تخزين المفاتيح، انتهاء مهلة المصافحة (30 ثانية)، ومنع التراجع.",
        "HandshakeHardeningTest (8/8): التوقيع الرقمي المرتبط بالسياق ومطابقة بصمات TOFU.",
        "اختبارات النماذج العامة (3/3): سلامة كبسولات البيانات وتراكيب الرسائل."
    ],
    border_color=COLOR_GREEN_TEXT
)

add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1),
    "⚠️ ما لم يتم التحقق منه ميدانياً بعد:",
    [
        "الاتصال الميداني بين هاتفين حقيقيين: لم يُجرَ اختبار شبكي بين جهازين فيزيائيين عبر راوتر Wi-Fi فعلي.",
        "اختبارات Android Instrumented: لم تُنفذ على محاكي أندرويد لعدم توفر بيئة عتادية في مرحلة الفحص.",
        "تكامل العتاد الحقيقي: التحقق من تخزين Keystore داخل معالج TEE يتطلب جهازاً فيزيائياً متاحاً.",
        "التغطية الشاملة لجميع المتطلبات: الاختبارات غطت وحدات التشفير والمصافحة فقط وليس كافة الـ 31 متطلباً.",
        "مصفوفة التتبع: موثقة بالكامل في docs/TraceabilityMatrix.md لتوضيح حالة كل متطلب."
    ],
    border_color=COLOR_AMBER_TEXT
)
set_speaker_notes(s11, "فصل حاسم بين نجاح الاختبارات وبين التحقق الميداني: نجاح 33 من 33 اختباراً يثبت صحة خوارزميات التشفير وتبادل المفاتيح وإدارة الجلسات في بيئة JVM. لكنه لا يعني إثبات عمل الشبكة بين جهازين أندرويد حقيقيين عبر Wi-Fi فعلي. توثيق هذا الفرق في مصفوفة TraceabilityMatrix.md هو ما يميز العمل الهندسي المنضبط.")

# ==============================================================================
# SLIDE 12: Roadmap & Gaps
# ==============================================================================
s12 = prs.slides.add_slide(blank_slide_layout)
add_header(s12, 12, "الأمانة الأكاديمية وخارطة الطريق", "الفجوات الأمنية المتبقية والتحسينات المستقبلية", "الاعتراف العلمي بالحدود الحالية وتحديد أولويات التطوير للمرحلة القادمة", "ROADMAP", "blue")

gap_w = Inches(2.78)
gap_h = Inches(3.6)

gaps = [
    ("1. قناة التحقق SAS", "مقارنة كود بصري أو صوتي (Short Auth String) أو مسح QR في أول لقاء لإغلاق ثغرة MITM الأولى نهائياً.", "NOT IMPLEMENTED"),
    ("2. حماية الـ Replay الكاملة", "إضافة أرقام تسلسلية (Sequence Numbers) مع نافذة منزلقة (Sliding Window) لرفض الحزم القديمة المكررة.", "NOT IMPLEMENTED"),
    ("3. فحص الملف الكامل", "حساب بصمة SHA-256 لكامل الملف بعد اكتمال تجميع كتل .part للتأكد من سلامة المحتوى الإجمالي.", "NOT IMPLEMENTED"),
    ("4. مفاتيح المجموعات", "تطبيق بروتوكول شجرة المفاتيح (Group Keys) لتشفير الغرف الجماعية ومسارات بث الوسائط بمفاتيح ديناميكية.", "NOT IMPLEMENTED")
]

for gidx, (gtitle, gdesc, gstat) in enumerate(gaps):
    gx = Inches(0.8 + gidx * 2.98)
    gy = Inches(1.8)
    add_card(s12, gx, gy, gap_w, gap_h, gtitle, [gdesc], border_color=COLOR_RED_TEXT)

# Final Conclusion Box
final_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.73), Inches(1.5))
final_box.fill.solid()
final_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
final_box.line.color.rgb = COLOR_ROYAL_BLUE
final_box.line.width = Pt(1.5)

tf_final = final_box.text_frame
tf_final.word_wrap = True
p_f_t = tf_final.paragraphs[0]
p_f_t.text = "🎓 الخلاصة الهندسية والتوصية الختامية:"
p_f_t.font.name = FONT_ARABIC
p_f_t.font.size = Pt(12)
p_f_t.font.bold = True
p_f_t.font.color.rgb = COLOR_PRIMARY_NAVY

p_f_d = tf_final.add_paragraph()
p_f_d.text = "تحول مشروع LOCAL CONTACT من تطبيق بسيط يعتمد على مفتاح ثابت مكشوف إلى نظام شبكي P2P مهيكل وفق أحدث مبادئ هندسة البرمجيات (SRS, Baseline, RTM) ومحصن بنواة تشفير معتمدة (AES-256-GCM, ECDH P-256, HKDF, Keystore, TOFU) مع مصفوفة تتبع صارمة تفصل بوضوح بين ما تم إثباته برمجياً وما يزال قيد التحسين المستقبلي."
p_f_d.font.name = FONT_ARABIC
p_f_d.font.size = Pt(10.5)
p_f_d.font.color.rgb = COLOR_TEXT_DARK
p_f_d.space_before = Pt(4)

set_speaker_notes(s12, "اختم العرض بقوة وبأمانة علمية تامة: لخّص الفجوات المتبقية والموثقة كخارطة طريق مستقبلية: 1- التحقق اليدوي SAS لمنع MITM في أول اتصال. 2- Replay Protection كامل بأرقام تسلسلية للحزم. 3- فحص سلامة الملف كاملاً SHA-256. 4- Group Keys للغرف. اختم بالتأكيد على أن المشروع انتقل من كود أولي بمفتاح ثابت إلى بنية هندسية رصينة موثقة ومختبرة وقابلة للتطوير المستمر.")

# Save presentation
out_pptx_en = os.path.abspath("LOCAL_CONTACT_Security_Presentation.pptx")
prs.save(out_pptx_en)
out_pptx_ar = os.path.join(os.path.dirname(out_pptx_en), "مناقشة_خوارزميات_التشفير_LOCAL_CONTACT.pptx")
import shutil
shutil.copyfile(out_pptx_en, out_pptx_ar)

print(f"Presentation saved successfully:")
print(f"1: {out_pptx_en} ({os.path.getsize(out_pptx_en):,} bytes)")
print(f"2: {out_pptx_ar} ({os.path.getsize(out_pptx_ar):,} bytes)")

